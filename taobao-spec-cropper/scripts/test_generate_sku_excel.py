"""使用临时商品目录验证 SKU 输出。"""
import json
import subprocess
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from generate_sku_excel import generate


class SkuTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.template = self.root / 'SKU模板.xls'
        wb = Workbook()
        wb.active.title = 'AI批量上传'
        wb.active.append(['颜色分类', '价格', '数量'])
        wb.active.append(['示例', 1, 1])
        wb.create_sheet('保留').sheet_state = 'hidden'
        wb.save(self.template)

    def product(self, name='商品', price='25.75'):
        folder = self.root / name
        folder.mkdir()
        (folder / '商品信息.json').write_text(json.dumps({'price': price}), encoding='utf-8-sig')
        specs = folder / '详情图_规格图_800_PNG'
        specs.mkdir()
        for name in ['02_浅蓝色.png', '01_白色.png']:
            (specs / name).touch()
        return folder

    def test_rows_template_and_no_overwrite(self):
        folder = self.product()
        output = generate(self.root, self.template, set())[0]
        self.assertEqual(output, folder / '详情图_规格图_800_PNG' / 'sku.xls')
        before = output.read_bytes()
        wb = load_workbook(BytesIO(before))
        self.assertEqual(list(wb.active.values), [('颜色分类', '价格', '数量'), ('01#白色 (半米价)', 25.75, 100), ('02#浅蓝色 (半米价)', 25.75, 100)])
        self.assertEqual(wb['保留'].sheet_state, 'hidden')
        self.assertEqual(wb.active['B2'].number_format, 'General')
        wb.close()
        with self.assertRaises(FileExistsError):
            generate(self.root, self.template, set())
        self.assertEqual(output.read_bytes(), before)

    def test_invalid_batch_writes_nothing(self):
        self.product('A')
        self.product('B', 'NaN')
        with self.assertRaises(ValueError):
            generate(self.root, self.template, set())
        self.assertFalse(list(self.root.glob('*/*/sku.xls')))
        self.assertEqual(len(generate(self.root, self.template, {'B'})), 1)

    def test_duplicate_number(self):
        folder = self.product()
        (folder / '详情图_规格图_800_PNG' / '01_重复.png').touch()
        with self.assertRaises(ValueError):
            generate(self.root, self.template, set())

    def test_opt_in_required(self):
        self.product()
        result = subprocess.run([sys.executable, str(Path(__file__).with_name('generate_sku_excel.py')), '--root', str(self.root), '--template', str(self.template)], capture_output=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(list(self.root.glob('*/*/sku.xls')))

    def test_single_product_spec_dir_and_custom_xls_name(self):
        folder = self.product()
        old = folder / '详情图_规格图_800_PNG'
        new = folder / '规格图'
        old.rename(new)
        output = generate(self.root, self.template, set(), product_dir=folder,
                          spec_dir_name='规格图', output_name='sku.xls',
                          output_in_spec_dir=True)[0]
        self.assertEqual(output.name, 'sku.xls')
        self.assertEqual(output.parent.name, '规格图')
        self.assertTrue(output.is_file())
        wb = load_workbook(BytesIO(output.read_bytes()))
        self.assertEqual(wb.active['A2'].value, '01#白色 (半米价)')
        self.assertEqual(wb.active['B2'].value, 25.75)
        wb.close()

    def test_output_in_product_dir_opt_out(self):
        folder = self.product()
        output = generate(self.root, self.template, set(), product_dir=folder,
                          output_name='商品_SKU.xlsx', output_in_spec_dir=False)[0]
        self.assertEqual(output, folder / '商品_SKU.xlsx')
        self.assertTrue(output.is_file())


if __name__ == '__main__':
    unittest.main()
