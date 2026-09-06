"""按商品信息和规格图文件名生成 SKU Excel。"""
from __future__ import annotations

import argparse
import json
import re
import math
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


def generate(root: Path, template: Path, excludes: set[str]) -> list[Path]:
    template_bytes = template.read_bytes()
    try:
        workbook = load_workbook(BytesIO(template_bytes))
    except Exception as exc:
        raise ValueError("模板必须是 XLSX 工作簿；旧版二进制 XLS 请先另存为 XLSX") from exc
    sheet_name = "AI批量上传" if "AI批量上传" in workbook.sheetnames else workbook.active.title
    if [workbook[sheet_name].cell(1, col).value for col in range(1, 4)] != ["颜色分类", "价格", "数量"]:
        raise ValueError("模板首行必须依次为：颜色分类、价格、数量")
    workbook.close()
    plans = []
    outputs: list[Path] = []
    for folder in sorted(p for p in root.iterdir() if p.is_dir() and (p / "商品信息.json").is_file() and p.name not in excludes):
        info = json.loads((folder / "商品信息.json").read_text(encoding="utf-8-sig"))
        try:
            price = float(info["price"])
            if isinstance(info["price"], bool) or not math.isfinite(price) or price < 0:
                raise ValueError("价格必须是非负有限数值")
        except (KeyError, TypeError, ValueError) as exc:
            raise ValueError(f"{folder.name}: 商品信息.json 的 price 无效") from exc
        spec_dirs = sorted(p for p in folder.iterdir() if p.is_dir() and "规格图_800_PNG" in p.name)
        if len(spec_dirs) != 1:
            raise ValueError(f"{folder.name}: 需要且只能有一个规格图目录")
        names = sorted((p.name for p in spec_dirs[0].iterdir() if p.is_file() and re.fullmatch(r"\d+_.+\.png", p.name, re.I)), key=lambda n: int(n.split("_", 1)[0]))
        if not names:
            raise ValueError(f"{folder.name}: 未找到规格图 PNG")
        rows = []
        for index, name in enumerate(names, 1):
            number, color = name[:-4].split("_", 1)
            if int(number) != index:
                raise ValueError(f"{folder.name}: 编号不连续，发现 {number}")
            rows.append([f"{number}#{color} (半米价)", price, 100])
        output = folder / f"{folder.name}_SKU.xlsx"
        if output.exists():
            raise FileExistsError(f"文件已存在，不覆盖：{output}")
        plans.append((folder, output, rows))
    if not plans:
        raise ValueError("未找到包含商品信息.json 的商品文件夹")
    for folder, output, rows in plans:
        workbook = load_workbook(BytesIO(template_bytes))
        sheet = workbook[sheet_name]
        sheet.delete_rows(2, max(sheet.max_row - 1, 1))
        for row in rows:
            sheet.append(row)
        sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width or 0, 36)
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=2, max_col=3):
            for cell in row:
                cell.number_format = "0" if cell.value == int(cell.value) else "General"
        with output.open("xb") as stream:
            workbook.save(stream)
        workbook.close()
        check = load_workbook(output, read_only=True)
        if list(check[sheet_name].iter_rows(min_row=2, max_col=3, values_only=True)) != [tuple(row) for row in rows]:
            raise ValueError(f"成品校验失败：{output}")
        check.close()
        outputs.append(output)
        print(f"{folder.name}：{len(rows)} 行，已生成")
    return outputs


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", required=True, type=Path)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--sku-excel", action="store_true", help="确认生成 SKU Excel")
    parser.add_argument("--exclude", action="append", default=[], help="排除商品文件夹，可重复")
    args = parser.parse_args()
    if not args.sku_excel:
        parser.error("需要显式传入 --sku-excel 才会生成文件")
    try:
        generate(args.root.resolve(), args.template.resolve(), set(args.exclude))
    except (ValueError, OSError) as exc:
        parser.exit(2, f"错误：{exc}\n")


if __name__ == "__main__":
    main()
